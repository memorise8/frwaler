# -*- coding: utf-8 -*-
"""최종 케파 병합 + 보고서(xlsx/md) 생성.

입력(있는 것만 사용):
  - baseline(측정 199): coverage_report.csv ∪ custom_crawler_totals.SNAPSHOT.csv
    (HAL 글로벌 >1.5M 제외)
  - count_only_totals.csv (초기 6분 count)
  - count_only_capped_1h.csv (1시간 재측정)
  - count_only_uncapped.csv (내부캡 무력화 재측정) — 있으면 최우선
  - per-site 로그로 count_status 분류
  - DB: collected, pdf_size_bytes

우선순위(사이트별 최종 source_total): uncapped > 1h > 초기 count > baseline
"""
import csv, glob, os, re, sqlite3
from collections import defaultdict

ROOT = "/data_raid/ruci_workspace/frwaler_job"
AUD = f"{ROOT}/scripts/audit"
LOGDIR = f"{AUD}/out/count_only_logs"
DB = f"{ROOT}/libertree-app/data/libertree.db"
MB_PER_DOC = 3.83
GLOBAL_HAL = 1_500_000

def num(x):
    try: return int(float(x))
    except: return None

def load_csv(path):
    return list(csv.DictReader(open(path))) if os.path.exists(path) else []

# ---- DB: collected + held bytes ----
c = sqlite3.connect(f"file:{DB}?mode=ro", uri=True)
collected = dict(c.execute(
    "SELECT s.site_id, COUNT(d.seq_id) FROM sites s LEFT JOIN documents d ON d.site_id=s.site_id GROUP BY s.site_id").fetchall())
sheet = dict(c.execute("SELECT site_id, sheet FROM sites").fetchall())
name = dict(c.execute("SELECT site_id, site_name FROM sites").fetchall())
held_bytes = dict(c.execute(
    "SELECT site_id, COALESCE(SUM(pdf_size_bytes),0) FROM documents GROUP BY site_id").fetchall())
all_sites = set(sheet) | set(collected)

# ---- baseline measured (HAL 글로벌 제외) ----
measured = {}
for path in ["coverage_report.csv", "custom_crawler_totals.SNAPSHOT.csv", "custom_crawler_totals.csv"]:
    for r in load_csv(f"{AUD}/{path}"):
        s = num(r.get("source_total")); m = r.get("method", "")
        if s and s > 0:
            if m == "hal_solr" and s > GLOBAL_HAL:
                continue
            measured[r["site_id"]] = (s, m)

# ---- count passes (completed 플래그 포함) ----
def load_count(path):
    out = {}
    for r in load_csv(f"{AUD}/{path}"):
        s = num(r.get("counted"))
        if s is not None:
            st = num(r.get("server_total")) or 0
            out[r["site_id"]] = (s, (r.get("completed", "").lower() == "true"),
                                 r.get("method", "count_crawl"), st)
    return out

cnt_initial = load_count("count_only_totals.csv")
cnt_1h = load_count("count_only_capped_1h.csv")
cnt_uncapped = load_count("count_only_uncapped.csv")
cnt_totalscan = load_count("count_only_totalscan.csv")

# 전역 server_total: 모든 패스에서 포착한 서버 총량 중 사이트별 최대
server_totals = {}
for _d in (cnt_initial, cnt_1h, cnt_uncapped, cnt_totalscan):
    for _sid, (_cnt, _comp, _meth, _st) in _d.items():
        if _st and _st > server_totals.get(_sid, 0):
            server_totals[_sid] = _st

# ---- count_status: 로그 스캔 ----
CAP_RE = re.compile(r"safety cap|MAX_PAGES|budget reached|wall-clock|budget exceeded", re.I)
END_RE = re.compile(r"no more items|empty page|keine Publikationen|no more|Done\.|not_?found", re.I)
def status_from_log(sid, completed):
    logs = glob.glob(f"{LOGDIR}/{sid}.log")
    if not logs:
        return "no_log"
    txt = open(logs[0], errors="replace").read()
    if not completed:
        return "external_timeout"
    if CAP_RE.search(txt):
        return "crawler_cap"
    if END_RE.search(txt):
        return "natural"
    return "unknown"

# ---- 사이트별 최종 병합 ----
rows = []
for sid in sorted(all_sites):
    src = None; method = ""; completed = None; origin = ""; server_total = 0
    if sid in cnt_uncapped:
        src, completed, method, server_total = cnt_uncapped[sid]; origin = "uncapped"
    elif sid in cnt_1h:
        src, completed, method, server_total = cnt_1h[sid]; origin = "1h"
    elif sid in cnt_initial:
        src, completed, method, server_total = cnt_initial[sid]; origin = "initial_count"
    elif sid in measured:
        src, method = measured[sid]; completed = True; origin = "api_measured"
    # count_status
    if origin == "api_measured":
        cstatus = "measured_api"
    elif src is None:
        cstatus = "unmeasured"
    else:
        cstatus = status_from_log(sid, bool(completed))
    # 서버가 알려준 실제 총량(모든 패스 중 최대)이 count보다 크면 정확값으로 채택
    server_total = max(server_total or 0, server_totals.get(sid, 0))
    if server_total and server_total > (src or 0):
        src = server_total
        cstatus = "server_total"  # API가 알려준 정확한 전량
    rows.append({
        "site_id": sid, "sheet": sheet.get(sid, ""), "site_name": name.get(sid, ""),
        "collected": collected.get(sid, 0), "held_bytes": held_bytes.get(sid, 0),
        "source_total": src, "count_status": cstatus, "origin": origin, "method": method,
    })

# ---- 집계 ----
tot_collected = sum(r["collected"] for r in rows)
tot_held_bytes = sum(r["held_bytes"] for r in rows)
meas_rows = [r for r in rows if r["source_total"] is not None]
tot_src = sum(r["source_total"] for r in meas_rows)
exact_rows = [r for r in meas_rows if r["count_status"] in ("natural", "measured_api", "server_total")]
under_rows = [r for r in meas_rows if r["count_status"] in ("crawler_cap", "external_timeout", "unknown", "no_log")]
tot_src_exact = sum(r["source_total"] for r in exact_rows)
tot_src_under = sum(r["source_total"] for r in under_rows)
unmeasured = [r for r in rows if r["source_total"] is None]

from collections import Counter
cstat = Counter(r["count_status"] for r in rows)

print("="*66)
print("최종 케파 집계")
print("="*66)
print(f"총 사이트: {len(rows)}")
print(f"보유(수집): {tot_collected:,}건 / held PDF {tot_held_bytes/1024/1024/1024:.1f} GB")
print(f"측정된 사이트: {len(meas_rows)}  (미측정 {len(unmeasured)})")
print(f"  가용(소스) 합계: {tot_src:,}건  → 커버리지 {100*tot_collected/tot_src:.1f}% (보유 대비)")
print(f"    ├ exact(자연종료+API): {tot_src_exact:,}  ({len(exact_rows)}개)")
print(f"    └ 과소치(내부캡+외부상한): {tot_src_under:,}  ({len(under_rows)}개) — 실제 더 큼")
print(f"전량수집 필요용량(하한 {tot_src:,}): ≈ {tot_src*MB_PER_DOC/1024/1024:.0f} TB")
print(f"\ncount_status 분포: {dict(cstat)}")

# ---- xlsx ----
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "요약"
    def put(r, vals):
        for i, v in enumerate(vals, 1): ws.cell(row=r, column=i, value=v)
    ws.append(["케파 최종 보고서 (count-only, 무저장)"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["보유 문서", tot_collected])
    ws.append(["보유 PDF 용량(GB)", round(tot_held_bytes/1024/1024/1024, 1)])
    ws.append(["실디스크(참고)", "libertree 1.7TB + data 32GB + DB 5.6GB ≈ 1.8TB"])
    ws.append([])
    ws.append(["측정된 사이트", len(meas_rows), f"/ 전체 {len(rows)}"])
    ws.append(["가용 총계(하한)", tot_src])
    ws.append(["  exact분", tot_src_exact])
    ws.append(["  과소치분(실제 더 큼)", tot_src_under])
    ws.append(["미측정 사이트", len(unmeasured)])
    ws.append(["전량수집 필요용량(TB, 하한)", round(tot_src*MB_PER_DOC/1024/1024, 1)])
    ws.append(["현재 디스크 여유(TB)", 3.6])
    ws.append([])
    ws.append(["※ 과소치 원인 3가지: ①외부상한 ②크롤러 내부캡 ③자동측정 불가. '가용 총계'는 하한이며 실제 전량은 더 큼."])

    ws2 = wb.create_sheet("사이트별")
    ws2.append(["site_id", "sheet", "site_name", "collected", "source_total", "count_status", "origin", "method"])
    for cell in ws2[1]: cell.font = Font(bold=True)
    for r in sorted(rows, key=lambda x: -(x["source_total"] or -1)):
        ws2.append([r["site_id"], r["sheet"], r["site_name"], r["collected"],
                    r["source_total"], r["count_status"], r["origin"], r["method"]])

    ws3 = wb.create_sheet("국가별_롤업")
    ws3.append(["sheet", "사이트수", "측정수", "collected", "source_total(측정분)", "커버리지%"])
    for cell in ws3[1]: cell.font = Font(bold=True)
    by = defaultdict(lambda: [0, 0, 0, 0])
    for r in rows:
        b = by[r["sheet"]]; b[0]+=1
        if r["source_total"] is not None: b[1]+=1; b[3]+=r["source_total"]
        b[2]+=r["collected"]
    for sh, b in sorted(by.items(), key=lambda x:-x[1][3]):
        cov = round(100*b[2]/b[3], 1) if b[3] else ""
        ws3.append([sh, b[0], b[1], b[2], b[3], cov])

    ws4 = wb.create_sheet("크롤러캡_상향필요")
    ws4.append(["site_id", "sheet", "collected", "source_total(과소)", "count_status", "note"])
    for cell in ws4[1]: cell.font = Font(bold=True)
    for r in sorted([x for x in rows if x["count_status"] == "crawler_cap"], key=lambda x:-(x["source_total"] or 0)):
        ws4.append([r["site_id"], r["sheet"], r["collected"], r["source_total"], r["count_status"],
                    "크롤러 내부캡으로 잘림 — 전량수집 시 캡 상향 필요"])

    ws5 = wb.create_sheet("상한재측정_대상")
    ws5.append(["site_id", "sheet", "collected", "source_total(하한)", "count_status"])
    for cell in ws5[1]: cell.font = Font(bold=True)
    for r in sorted([x for x in rows if x["count_status"] == "external_timeout"], key=lambda x:-(x["source_total"] or 0)):
        ws5.append([r["site_id"], r["sheet"], r["collected"], r["source_total"], r["count_status"]])

    out = f"{AUD}/capacity_final_report.xlsx"
    wb.save(out)
    print(f"\n-> {out} 저장 (시트 5개)")
except ImportError:
    print("openpyxl 없음 — xlsx 스킵")

# ---- md 요약 ----
with open(f"{AUD}/capacity_final_report.md", "w", encoding="utf-8") as f:
    f.write(f"# 케파 최종 보고\n\n")
    f.write(f"- 보유: {tot_collected:,}건 / {tot_held_bytes/1024/1024/1024:.1f}GB (실디스크 ≈1.8TB)\n")
    f.write(f"- 측정 {len(meas_rows)}/{len(rows)} 사이트, 가용 하한 **{tot_src:,}건**\n")
    f.write(f"  - exact {tot_src_exact:,} / 과소치 {tot_src_under:,} (실제 더 큼)\n")
    f.write(f"- 전량수집 필요용량(하한): **≈{tot_src*MB_PER_DOC/1024/1024:.0f}TB** (여유 3.6TB)\n")
    f.write(f"- count_status: {dict(cstat)}\n")
print(f"-> {AUD}/capacity_final_report.md 저장")
