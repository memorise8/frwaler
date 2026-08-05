# -*- coding: utf-8 -*-
"""exact_probe 전체 544개 — **청크 배치판** (40개/프로세스, FD누수 격리).

배경: 단일 프로세스로 544개 연속 probe 시 리소스 고갈로 fetch 전멸(2026-08-04).
각 청크를 별도 subprocess로 실행해 누수를 격리하고, --resume 으로 확보된 사이트는 건너뜀.
완료 후 tests.xlsx 'exact_probe_실패' 시트 재작성 (기존 시트 보존).
"""
import csv, os, subprocess, sys, time
from pathlib import Path

ROOT = Path("/data_raid/ruci_workspace/frwaler_job")
AUD = ROOT / "scripts/audit"
PY = str(ROOT / ".venv/bin/python")
ENV = dict(os.environ, PYTHONPATH=str(ROOT))
OUT = AUD / "exact_probe.csv"
CHUNK = 40

def log(m):
    print(f"[exact-chunked {time.strftime('%H:%M:%S')}] {m}", flush=True)

def has_value(row):
    return str(row.get("exact_total", "")).strip().isdigit()

def load_out():
    if not OUT.exists():
        return {}
    return {r["site_id"]: r for r in csv.DictReader(open(OUT, encoding="utf-8", newline=""))}

ids = []
with open(ROOT / "html_list.csv", encoding="utf-8-sig", newline="") as f:
    for r in csv.DictReader(f):
        sid = (r.get("site_id") or "").strip()
        if sid:
            ids.append(sid)

existing = load_out()
todo = [s for s in ids if not (s in existing and has_value(existing[s]))]
log(f"전체 {len(ids)} / 확보 {len(ids)-len(todo)} / 재시도 {len(todo)} (청크 {CHUNK})")

for i in range(0, len(todo), CHUNK):
    chunk = todo[i:i+CHUNK]
    log(f"청크 {i//CHUNK+1}/{(len(todo)+CHUNK-1)//CHUNK}: {len(chunk)}개")
    subprocess.run([PY, str(AUD / "exact_probe.py"),
                    "--only", *chunk,
                    "--workers", "6", "--walk-wall", "900", "--resume",
                    "--out-csv", str(OUT)], env=ENV)
    # 청크별 성공률 로그
    cur = load_out()
    got = sum(1 for s in chunk if s in cur and has_value(cur[s]))
    log(f"  청크 결과: {got}/{len(chunk)} 확보")

# ---- tests.xlsx 실패사유 시트 재작성 ----
REASON = {
    "no_crawler":       ("실패", "크롤러가 등록돼 있지 않음"),
    "instantiate_fail": ("실패", "크롤러 초기화 실패 (코드 오류)"),
    "capture_fail":     ("실패", "크롤러가 요청을 보내지 못함 (즉시 에러 — 사이트 차단/크롤러 고장)"),
    "no_list_url":      ("실패", "리스트 URL/페이지 파라미터 식별 불가 (1페이지에 page 파라미터 미노출)"),
    "fetch_fail":       ("실패", "리스트 페이지 요청 실패 (차단·타임아웃)"),
    "no_per_page":      ("실패", "리스트에서 게시물 링크 패턴 식별 불가 (JS 렌더링/특수 구조)"),
    "none":             ("실패", "총량 신호 없음 (총건수 텍스트·페이지네이션 모두 부재)"),
    "error":            ("실패", "probe 내부 오류"),
    "list_walk_partial": ("부분", "리스트 walk 시간상한 도달 — 값은 하한(실제 더 큼)"),
    "total_text_only":  ("성공(참고)", "게시물 링크 미식별이나 총건수 텍스트로 확보"),
}
rows = list(load_out().values())
ok = [r for r in rows if r["method"] in ("total_text", "lastpage_verified", "list_walk", "total_text_only")]
bad = [r for r in rows if r["method"] not in ("total_text", "lastpage_verified", "list_walk", "total_text_only")]
log(f"probe 최종: 정확값 {len(ok)} / 실패·부분 {len(bad)}")

import openpyxl
from openpyxl.styles import Font
XLSX = ROOT / "tests.xlsx"
wb = openpyxl.load_workbook(XLSX)
SHEET = "exact_probe_실패"
if SHEET in wb.sheetnames:
    del wb[SHEET]
ws = wb.create_sheet(SHEET)
ws.append(["site_id", "구분", "왜 안되는가(사유)", "probe method", "상세 신호", "비고"])
for c in ws[1]:
    c.font = Font(bold=True)
order = {"실패": 0, "부분": 1, "성공(참고)": 2}
listed = sorted(
    [r for r in rows if r["method"] in REASON],
    key=lambda r: (order[REASON[r["method"]][0]], r["site_id"]))
for r in listed:
    cat, why = REASON[r["method"]]
    ws.append([r["site_id"], cat, why, r["method"], r.get("signal", ""), (r.get("note", "") or "")[:180]])
ws.append([])
ws.append([f"전체 probe {len(rows)}개 중 정확값 확보 {len(ok)}개 / 실패·부분 {len(bad)}개 "
           f"(작성 {time.strftime('%Y-%m-%d %H:%M')})"])
for col, w in zip("ABCDEF", [34, 10, 58, 18, 40, 50]):
    ws.column_dimensions[col].width = w
wb.save(XLSX)
log(f"tests.xlsx '{SHEET}' 재작성 완료 (실패·부분 {len(bad)}개)")
log("=== 완료 ===")
