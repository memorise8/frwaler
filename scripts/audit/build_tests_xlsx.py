# -*- coding: utf-8 -*-
"""test.xlsx 각 URL의 현재 수집 현황을 판정해 tests.xlsx로 작성.

- 신규 크롤러(gob.mx, HRB, datos 등) 포함 현재 libertree.db 기준.
- 호스트(도메인) 단위 대조: 해당 호스트에서 수집된 문서가 1건 이상이면 '수집'.
- 시트1 'URL수집현황': 원본 1995행 + 호스트/수집여부/호스트문서수/매칭site_id
- 시트2 '요약': 시트(국가)별 + 전체 수집률
"""
import sqlite3
from urllib.parse import urlparse
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = "/data_raid/ruci_workspace/frwaler_job"
DB = f"{ROOT}/libertree-app/data/libertree.db"

def host(u):
    try:
        h = urlparse(str(u).strip()).netloc.lower()
        return h[4:] if h.startswith("www.") else h
    except Exception:
        return ""

c = sqlite3.connect(f"file:{DB}?mode=ro", uri=True)
counts = dict(c.execute("SELECT site_id,count(*) FROM documents GROUP BY site_id").fetchall())
host_docs = defaultdict(int)
host_sids = defaultdict(list)
for sid, url in c.execute("SELECT site_id, site_url FROM sites"):
    h = host(url)
    host_docs[h] += counts.get(sid, 0)
    if counts.get(sid, 0) > 0:
        host_sids[h].append(sid)

src = openpyxl.load_workbook(f"{ROOT}/test.xlsx", read_only=True)
rows = [r for r in src["Sheet1"].iter_rows(values_only=True)]
header = rows[0]
data = rows[1:]

out = openpyxl.Workbook()
ws = out.active
ws.title = "URL수집현황"
hdr = list(header) + ["호스트", "수집여부", "호스트문서수", "매칭_site_id"]
ws.append(hdr)
bold = Font(bold=True)
fill_hdr = PatternFill("solid", fgColor="1F4E78")
white = Font(bold=True, color="FFFFFF")
for i, _ in enumerate(hdr, 1):
    c1 = ws.cell(row=1, column=i); c1.font = white; c1.fill = fill_hdr

by_sheet = defaultdict(lambda: [0, 0])
col = tot = 0
fill_no = PatternFill("solid", fgColor="F8CBAD")
for r in data:
    sheet, seq, org, url, cat, sub = r
    if not url:
        continue
    tot += 1
    h = host(url)
    n = host_docs.get(h, 0)
    ok = n > 0
    by_sheet[sheet][1] += 1
    if ok:
        col += 1; by_sheet[sheet][0] += 1
    sids = ",".join(host_sids.get(h, [])[:3])
    ws.append(list(r) + [h, "수집" if ok else "미수집", n, sids])
    if not ok:
        ws.cell(row=ws.max_row, column=hdr.index("수집여부") + 1).fill = fill_no

# column widths
widths = [16, 8, 26, 60, 16, 10, 26, 8, 12, 30]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.freeze_panes = "A2"

# --- 요약 시트 ---
s2 = out.create_sheet("요약")
s2.append(["test.xlsx 최종 수집률 리포트 (신규 크롤러 포함, 호스트 기준)"])
s2.append([f"기준일 DB / 총 사이트 {c.execute('SELECT COUNT(*) FROM sites').fetchone()[0]}개 / 총 문서 {sum(counts.values()):,}건"])
s2.append([])
s2.append(["전체 URL", tot])
s2.append(["수집", col, f"{100*col/tot:.1f}%"])
s2.append(["미수집", tot - col])
s2.append([])
s2.append(["시트(국가/그룹)", "수집", "전체", "수집률%"])
for i in range(1, 5):
    s2.cell(row=8, column=i).font = bold
for sheet, (cc, t) in sorted(by_sheet.items(), key=lambda x: (x[1][0]/x[1][1] if x[1][1] else 0)):
    s2.append([sheet, cc, t, round(100*cc/t, 0)])
s2.cell(row=1, column=1).font = Font(bold=True, size=13)
s2.column_dimensions["A"].width = 34
for col_ in "BCD":
    s2.column_dimensions[col_].width = 12

out.save(f"{ROOT}/tests.xlsx")
print(f"tests.xlsx 저장 완료")
print(f"수집률: {col}/{tot} = {100*col/tot:.1f}%  (미수집 {tot-col})")
