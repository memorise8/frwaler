# -*- coding: utf-8 -*-
"""sweep 완료 대기 → E-15 잔여 sweep → **전체 544개 exact_probe** → 실패사유를 tests.xlsx에 기록.

사용자 지시(2026-08-04): "지금 작업 끝나면 전부 exact_probe 해줘,
exact_probe가 안되면 왜 안되는지 tests.xlsx에 작성해줘"

단계:
  1) 현재 run_c_full.py / harness 종료 대기 (/proc 스캔)
  2) run_c_full.py 1회 실행 — 재개 로직이 잔여(E-15 포함) 자동 측정
  3) exact_probe --only <html_list 544 전부> --resume  (이미 확보된 사이트 건너뜀)
  4) tests.xlsx에 'exact_probe_실패' 시트 추가(기존 시트 보존) — 실패/부분 사이트와 사유
전 과정 count-only(무저장), DB 읽기전용.
"""
import csv, os, subprocess, sys, time
from pathlib import Path

ROOT = Path("/data_raid/ruci_workspace/frwaler_job")
AUD = ROOT / "scripts/audit"
PY = str(ROOT / ".venv/bin/python")
ENV = dict(os.environ, PYTHONPATH=str(ROOT))

def log(m):
    print(f"[c-then-exact {time.strftime('%H:%M:%S')}] {m}", flush=True)

def running(marker):
    me = os.getpid()
    for pid in os.listdir("/proc"):
        if not pid.isdigit() or int(pid) == me:
            continue
        try:
            cl = open(f"/proc/{pid}/cmdline", "rb").read().replace(b"\0", b" ").decode(errors="replace")
        except Exception:
            continue
        if marker in cl and "run_c_then_exact_all" not in cl:
            return True
    return False

def wait_gone(marker):
    while running(marker):
        time.sleep(60)

# ---- 1) 현재 sweep 종료 대기 ----
log("현재 sweep 종료 대기...")
wait_gone("run_c_full.py")
wait_gone("count_only_harness.py --only")
log("sweep 종료 확인")
time.sleep(5)

# ---- 2) 잔여(E-15 등) sweep ----
log("run_c_full.py 재실행 (잔여 자동 측정)")
subprocess.run([PY, str(AUD / "run_c_full.py")], env=ENV)
log("잔여 sweep 완료")

# ---- 3) 전체 544개 exact_probe ----
ids = []
with open(ROOT / "html_list.csv", encoding="utf-8-sig", newline="") as f:
    for r in csv.DictReader(f):
        sid = (r.get("site_id") or "").strip()
        if sid:
            ids.append(sid)
log(f"exact_probe 전체 {len(ids)}개 (resume, workers 6, walk_wall 900s)")
subprocess.run([PY, str(AUD / "exact_probe.py"),
                "--only", *ids,
                "--workers", "6", "--walk-wall", "900", "--resume",
                "--out-csv", str(AUD / "exact_probe.csv")], env=ENV)
log("exact_probe 완료")

# ---- 4) tests.xlsx 실패사유 시트 ----
REASON = {
    "no_crawler":       ("실패", "크롤러가 등록돼 있지 않음"),
    "instantiate_fail": ("실패", "크롤러 초기화 실패 (코드 오류)"),
    "capture_fail":     ("실패", "크롤러가 요청을 보내지 못함 (즉시 에러 — 사이트 차단/크롤러 고장)"),
    "no_list_url":      ("실패", "리스트 URL/페이지 파라미터 식별 불가 (1페이지에 page 파라미터 미노출)"),
    "fetch_fail":       ("실패", "리스트 페이지 요청 실패 (차단·타임아웃)"),
    "no_per_page":      ("실패", "리스트에서 게시물 링크 패턴 식별 불가 (JS 렌더링/특수 구조)"),
    "none":             ("실패", "총량 신호 없음 (총건수 텍스트·페이지네이션 모두 부재)"),
    "error":            ("실패", "probe 내부 오류"),
    "list_walk_partial": ("부분", "리스트 walk 시간상한 도달 — 값은 하한(실제 더 큼)"),
    "total_text_only":  ("성공(참고)", "게시물 링크 미식별이나 총건수 텍스트로 확보"),
}
rows = list(csv.DictReader(open(AUD / "exact_probe.csv", encoding="utf-8", newline="")))
ok = [r for r in rows if r["method"] in ("total_text", "lastpage_verified", "list_walk", "total_text_only")]
bad = [r for r in rows if r["method"] not in ("total_text", "lastpage_verified", "list_walk", "total_text_only")]
log(f"probe 결과: 성공 {len(ok)} / 실패·부분 {len(bad)}")

import openpyxl
from openpyxl.styles import Font
XLSX = ROOT / "tests.xlsx"
wb = openpyxl.load_workbook(XLSX)
SHEET = "exact_probe_실패"
if SHEET in wb.sheetnames:
    del wb[SHEET]
ws = wb.create_sheet(SHEET)
ws.append(["site_id", "구분", "왜 안되는가(사유)", "probe method", "상세 신호", "비고"])
for c in ws[1]:
    c.font = Font(bold=True)
order = {"실패": 0, "부분": 1, "성공(참고)": 2}
listed = sorted(
    [r for r in rows if r["method"] in REASON],
    key=lambda r: (order[REASON[r["method"]][0]], r["site_id"]))
for r in listed:
    cat, why = REASON[r["method"]]
    ws.append([r["site_id"], cat, why, r["method"], r.get("signal", ""), r.get("note", "")[:180]])
# 요약행
ws.append([])
ws.append([f"전체 probe {len(rows)}개 중 정확값 확보 {len(ok)}개 / 실패·부분 {len(bad)}개 "
           f"(작성 {time.strftime('%Y-%m-%d %H:%M')})"])
for col, w in zip("ABCDEF", [34, 10, 58, 18, 40, 50]):
    ws.column_dimensions[col].width = w
wb.save(XLSX)
log(f"tests.xlsx '{SHEET}' 시트 작성 완료 (실패·부분 {len(bad)}개 사유 기록)")
log("=== 전체 파이프라인 완료 ===")
