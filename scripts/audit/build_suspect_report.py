# -*- coding: utf-8 -*-
"""업체 확인 요청용 '의심 사이트' 리포트 + test.xlsx 수집현황 반영본 생성.

현재 libertree.db 기준으로 매칭을 다시 돌려:
  - test_수집현황반영_<date>.xlsx  (원본 6컬럼 + 데이터개수/수집건수/수집여부/매칭크롤러)
  - 의심사이트_업체확인_<date>.xlsx (요약 + 카테고리별 시트 + 업체확인질문)
수집이 끝날 때마다 재실행하면 최신본으로 갱신됨.
"""
from __future__ import annotations
import csv, sqlite3, collections
from urllib.parse import urlparse
from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path("/data_raid/ruci_workspace/frwaler_job")
OUT = ROOT / "docs/deliverables"
DATE = "20260802"

def dom(u):
    try: return urlparse(str(u)).netloc.lower().replace("www.", "").strip()
    except: return ""
def pth(u):
    try: return urlparse(str(u)).path.rstrip("/").lower()
    except: return ""
def numi(x):
    try: return int(float(x))
    except: return None

# --- DB (신규 크롤러 반영) ---
c = sqlite3.connect(f"file:{ROOT}/libertree-app/data/libertree.db?mode=ro", uri=True)
sites = c.execute("SELECT site_id,site_url FROM sites").fetchall()
collected = dict(c.execute("SELECT site_id,COUNT(*) FROM documents GROUP BY site_id").fetchall())
dom2sites = collections.defaultdict(list)
for sid, surl in sites:
    dom2sites[dom(surl)].append((sid, pth(surl)))
docdoms = collections.Counter()
for (mu,) in c.execute("SELECT meta_url FROM documents WHERE meta_url IS NOT NULL"):
    d = dom(mu)
    if d: docdoms[d] += 1

cap = {}
capp = ROOT / "scripts/audit/capacity_final.csv"
if capp.exists():
    for r in csv.DictReader(open(capp, encoding="utf-8-sig")):
        cap[r["site_id"]] = (numi(r.get("max_to_collect")), r.get("source", ""))

probe = {}
pp = ROOT / "scripts/audit/uncollected_probe.csv"
if pp.exists():
    for r in csv.DictReader(open(pp, encoding="utf-8-sig")):
        probe[r["URL주소"].strip()] = r

def match(u):
    d = dom(u); p = pth(u); cands = dom2sites.get(d, [])
    if cands:
        best = None; bl = -1
        for sid, sp in cands:
            if sp and (p.startswith(sp) or sp.startswith(p)):
                if len(sp) > bl: best = sid; bl = len(sp)
        if best is None: best = cands[0][0]
        col = collected.get(best, 0)
        mx = cap.get(best, (None, ""))[0]
        return best, col, mx, ("수집됨" if col > 0 else "크롤러有_미수집")
    if docdoms.get(d, 0) > 0:
        return "(사이트미등록)", docdoms[d], None, "수집됨(도메인매칭)"
    return "", 0, None, "미수집(크롤러없음)"

# --- 원본 test.xlsx ---
src = load_workbook(ROOT / "test.xlsx", read_only=True)["Sheet1"]
rows = list(src.iter_rows(values_only=True))
hdr = list(rows[0])
data = [r for r in rows[1:] if r and r[3] is not None]

hf = PatternFill("solid", fgColor="2F5496"); hfont = Font(color="FFFFFF", bold=True)
okf = PatternFill("solid", fgColor="E2EFDA"); nof = PatternFill("solid", fgColor="FCE4D6")
warnf = PatternFill("solid", fgColor="FFF2CC")

def style_header(ws, cols):
    for i in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=i); cell.fill = hf; cell.font = hfont
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"

# ============ 1) test 수집현황 반영본 ============
wb1 = Workbook(); ws = wb1.active; ws.title = "Sheet1"
nh = hdr + ["데이터개수(가용총량)", "수집건수", "수집여부", "매칭크롤러(site_id)"]
ws.append(nh); style_header(ws, nh)
matched = []
for r in data:
    sid, col, mx, st = match(r[3]); matched.append((r, sid, col, mx, st))
    ws.append(list(r) + [mx if mx is not None else "", col, st, sid])
    for i in range(1, len(nh) + 1):
        ws.cell(row=ws.max_row, column=i).fill = okf if st.startswith("수집") else nof
for i, w in enumerate([14, 8, 26, 50, 12, 13, 15, 9, 15, 28], 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.auto_filter.ref = f"A1:{get_column_letter(len(nh))}{ws.max_row}"
wb1.save(OUT / f"test_수집현황반영_{DATE}.xlsx")

# ============ 2) 의심 사이트 리포트 ============
wb2 = Workbook()

# 카테고리 계산
uncol = [(r, sid, col, mx, st) for (r, sid, col, mx, st) in matched if not st.startswith("수집")]
def pv(u): return probe.get(str(u).strip(), {})

cat_done_label = [(r, sid) for (r, sid, col, mx, st) in uncol if "완료" in str(r[0])]
cat_dead = [(r, sid) for (r, sid, col, mx, st) in uncol
            if pv(r[3]).get("크롤판정") == "죽음/접속불가" or str(pv(r[3]).get("상태코드", "")).startswith("ERR")
            or pv(r[3]).get("상태코드") in ("404", "410", "521")]
cat_blocked = [(r, sid) for (r, sid, col, mx, st) in uncol
               if pv(r[3]).get("크롤판정") == "차단(봇방어)" or pv(r[3]).get("상태코드") in ("403", "401")]
# 카운트 스코프 확인필요 (대량 + 근거가 html/onpage/전체카탈로그)
cat_scope = []
SCOPE_NOTE = {
    "e-stat-go-jp-stat-search": "통계표(Excel/CSV) 레코드 170만 — 논문 아님. '문서'에 데이터셋 포함 여부 확인",
    "ntrs-nasa-gov-search": "전체 카탈로그 총량. 크롤러 필터(회의논문 등)와 스코프 일치 여부 확인",
    "ga-gov-au-data-pubs": "GA 전체 eCat 카탈로그 수치. 요청 URL은 1970년대 부분집합 → 스코프 과다 가능",
    "dtic-dimensions-ai-discover": "25,000은 라운드넘버 정규식 추정치 → 실제값 재확인 필요",
    "sonar-ch-global": "필터(document_type) 적용 총량. 대상 범위 확인",
    "etera-ee-browse": "에스토니아 아카이브 전체. 자료 성격(문화유산 등) 확인",
    "gov-scot-publications": "온페이지 카운트 기반. 스코프 확인",
    "gov-uk-search": "gov.uk 전체 finder 총량. 대상 부처/유형 범위 확인",
}
seen = set()
for r, sid in [(r, sid) for (r, sid, col, mx, st) in matched]:
    if sid in seen: continue
    mxv, srcv = cap.get(sid, (None, ""))
    if mxv and mxv >= 20000 and ("html_count_regex" in srcv or srcv == "big_api"):
        cat_scope.append((sid, mxv, srcv, SCOPE_NOTE.get(sid, "대량 측정치 — 크롤러 스코프와 일치하는지 확인")))
        seen.add(sid)
cat_scope.sort(key=lambda x: -x[1])

# 요약 시트
ws = wb2.active; ws.title = "요약"
sumcols = ["의심 카테고리", "건수", "무엇을 확인해야 하나 (업체 확인 요청)"]
ws.append(sumcols); style_header(ws, sumcols)
summary = [
    ("① '완료' 라벨인데 미수집", len(cat_done_label),
     "업체가 '완료'로 표기한 URL이 실제로는 수집 안 됨 → 업체가 어디에/어떻게 수집했는지, 우리 대상과 동일한지 확인"),
    ("② URL 문제 (죽음/404/410)", len(cat_dead),
     "소스 목록의 URL이 사라짐/오류 → 업체에 최신 URL 요청 또는 목록에서 제외"),
    ("③ 접근 차단 (봇방어/CAPTCHA)", len(cat_blocked),
     "봇차단으로 접근 불가 → 업체가 어떻게 수집했는지(공식 API/제휴/수동?) 확인 또는 대상 제외"),
    ("④ 카운트 스코프 확인필요 (대량)", len(cat_scope),
     "총량이 크지만 '전체 카탈로그/통계표/추정치'라 크롤러 실제 대상과 다를 수 있음 → 스코프 정의 합의"),
]
for name, n, q in summary:
    ws.append([name, n, q]);
    for i in range(1, 4): ws.cell(row=ws.max_row, column=i).fill = warnf
for i, w in enumerate([30, 8, 90], 1): ws.column_dimensions[get_column_letter(i)].width = w
ws.append([]); ws.append(["※ 수집이 진행 중이면 ①②③은 완료 후 줄어들 수 있음. 완료 시 재생성 요망.", "", ""])

def add_url_sheet(title, items, note):
    ws = wb2.create_sheet(title)
    cols = ["시트명", "일련번호", "기관명", "URL주소", "구분", "매칭크롤러", "프로빙판정", "상태코드"]
    ws.append(cols); style_header(ws, cols)
    for r, sid in items:
        p = pv(r[3])
        ws.append([r[0], r[1], r[2], r[3], r[4], sid, p.get("크롤판정", ""), p.get("상태코드", "")])
        for i in range(1, len(cols) + 1): ws.cell(row=ws.max_row, column=i).fill = warnf
        ws.cell(row=ws.max_row, column=4).hyperlink = str(r[3]); ws.cell(row=ws.max_row, column=4).font = Font(color="0563C1", underline="single")
    for i, w in enumerate([14, 8, 26, 52, 12, 26, 20, 10], 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"

add_url_sheet("①완료라벨_미수집", cat_done_label, "")
add_url_sheet("②URL문제_죽음404", cat_dead, "")
add_url_sheet("③접근차단", cat_blocked, "")

ws = wb2.create_sheet("④카운트스코프확인")
cols = ["site_id", "측정총량", "측정근거", "업체 확인 요청 사항"]
ws.append(cols); style_header(ws, cols)
for sid, mxv, srcv, note in cat_scope:
    ws.append([sid, mxv, srcv, note])
    for i in range(1, len(cols) + 1): ws.cell(row=ws.max_row, column=i).fill = warnf
for i, w in enumerate([36, 12, 20, 70], 1): ws.column_dimensions[get_column_letter(i)].width = w

wb2.save(OUT / f"의심사이트_업체확인_{DATE}.xlsx")

# 콘솔 요약
ok = sum(1 for (_, _, _, _, st) in matched if st.startswith("수집"))
print(f"수집률: {ok}/{len(matched)} ({100*ok/len(matched):.1f}%)")
for name, n, q in summary: print(f"  {name}: {n}")
print(f"\n저장:\n  {OUT}/test_수집현황반영_{DATE}.xlsx\n  {OUT}/의심사이트_업체확인_{DATE}.xlsx")
