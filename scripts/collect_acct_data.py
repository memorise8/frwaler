"""acct_data.xlsx 링크 수집기 — KASB 5개 보드 + KICPA IFRS 실무사례.

기존 crawler.fino_acct 의 세션/파서/다운로더를 재사용한다.

  Phase A  xlsx 상세 URL(행사·교육자료 / 기고자료 / 회계기준적용의견서) → 제목+첨부
  Phase B  xlsx 목록 URL(List2003 / List2005 / List016008 / eduAccstdList) → A 미포함 첨부
  Phase C  KICPA IFRS 실무사례(list.brd → read.brd → fileMngr) 전건

manifest.ndjson 에 파일 단위로 append 하며, 재실행 시 이미 받은 키는 건너뛴다(멱등).

사용법:
    python -m scripts.collect_acct_data                 # 전체
    python -m scripts.collect_acct_data --phase A       # 특정 단계만
    python -m scripts.collect_acct_data --limit 5       # 스모크
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
import time
from pathlib import Path
from urllib.parse import parse_qs, urljoin, urlparse

import openpyxl
from bs4 import BeautifulSoup

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from crawler.fino_acct.fetch import download_attachment, new_session, safe_filename
from crawler.fino_acct.models import AttachmentLink
from crawler.fino_acct.parsers import clean_text, extract_links

REPO_ROOT = Path(__file__).resolve().parents[1]
XLSX_PATH = REPO_ROOT / "acct_data.xlsx"
OUT_ROOT = REPO_ROOT / "data" / "acct_data_20260805"
MANIFEST = OUT_ROOT / "manifest.ndjson"
DELAY = 0.5

KICPA_BOARD = "accstd02"
KICPA_LIST = "https://www.kicpa.or.kr/board/list.brd"
KICPA_READ = "https://www.kicpa.or.kr/board/read.brd"
KICPA_BLTN_RE = re.compile(r"readBulletin\('([^']+)'\s*,\s*'([^']+)'\)")

# 목록 페이지(첨부가 목록에 직접 노출됨). eduSmartList 는 YouTube 라 파일 없음 → 제외.
LIST_PAGES = {
    "List2003.do": "행사·교육자료",
    "List2005.do": "기고자료",
    "List016008.do": "정착지원TF",
    "eduAccstdList.do": "교육자료",
}
DETAIL_PAGES = ("View2003.do", "View2005.do", "opinionView.do")


# --------------------------------------------------------------------------- manifest


def load_manifest() -> tuple[list[dict], set[str]]:
    if not MANIFEST.exists():
        return [], set()
    rows: list[dict] = []
    with MANIFEST.open(encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if line:
                rows.append(json.loads(line))
    done = {r["key"] for r in rows if r.get("status") == "downloaded"}
    return rows, done


def append_manifest(record: dict) -> None:
    MANIFEST.parent.mkdir(parents=True, exist_ok=True)
    with MANIFEST.open("a", encoding="utf-8") as fh:
        _ = fh.write(json.dumps(record, ensure_ascii=False) + "\n")


def sha256_of(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


# --------------------------------------------------------------------------- xlsx


def read_xlsx_links() -> tuple[list[dict], list[dict], str | None]:
    """(상세 링크, 목록 링크, KICPA 진입 URL)."""
    workbook = openpyxl.load_workbook(XLSX_PATH)
    details: list[dict] = []
    lists: list[dict] = []
    kicpa_url: str | None = None

    sheet = workbook.worksheets[0]
    section = ""
    for row in sheet.iter_rows():
        if row[0].value:
            section = clean_text(str(row[0].value))
        title = clean_text(str(row[1].value)) if len(row) > 1 and row[1].value else ""
        for cell in row:
            if not cell.hyperlink:
                continue
            url = cell.hyperlink.target
            name = urlparse(url).path.rsplit("/", 1)[-1]
            entry = {"section": section, "title": title, "url": url, "row": cell.row}
            if name in LIST_PAGES:
                entry["section"] = entry["section"] or LIST_PAGES[name]
                lists.append(entry)
            elif name in DETAIL_PAGES:
                details.append(entry)

    for other in workbook.worksheets[1:]:
        for row in other.iter_rows():
            for cell in row:
                if cell.hyperlink and "kicpa.or.kr" in cell.hyperlink.target:
                    kicpa_url = cell.hyperlink.target

    # 동일 URL 중복 제거(시트에 같은 글이 두 번 걸린 경우)
    seen: set[str] = set()
    deduped: list[dict] = []
    for entry in details:
        if entry["url"] in seen:
            continue
        seen.add(entry["url"])
        deduped.append(entry)
    return deduped, lists, kicpa_url


# --------------------------------------------------------------------------- KASB


def kasb_key(link: AttachmentLink) -> str:
    return f"kasb:{link.post_file_no}:{link.post_file_seq}"


def collect_kasb_page(
    session,
    entry: dict,
    done: set[str],
    stats: dict,
    limit: int | None,
) -> None:
    url = entry["url"]
    try:
        response = session.get(url, timeout=30)
    except Exception as exc:  # noqa: BLE001 - 페이지 단위 실패는 기록만
        print(f"  [ERR] {url} :: {exc}", flush=True)
        stats["page_errors"] += 1
        return
    if response.status_code != 200:
        print(f"  [{response.status_code}] {url}", flush=True)
        stats["page_errors"] += 1
        return

    soup = BeautifulSoup(response.content, "html.parser")
    links = extract_links(url, soup)
    stats["pages"] += 1

    out_dir = OUT_ROOT / "kasb"
    for link in links.attachments:
        if not link.post_file_no:  # KASB 다운로드 패턴이 아닌 잡링크
            continue
        key = kasb_key(link)
        if key in done:
            stats["skipped"] += 1
            continue
        if limit is not None and stats["downloaded"] >= limit:
            return
        prefix = f"{link.post_file_no}_{link.post_file_seq}"
        result = download_attachment(session, link, out_dir, prefix, DELAY)
        record = {
            "key": key,
            "source": "kasb",
            "section": entry["section"],
            "title": entry["title"],
            "page_url": url,
            "file_no": link.post_file_no,
            "file_seq": link.post_file_seq,
            "filename": link.filename,
            "status": result.status,
            "content_type": result.content_type,
            "size_bytes": result.size_bytes,
        }
        if result.local_path is not None:
            record["path"] = str(result.local_path.relative_to(OUT_ROOT))
            record["sha256"] = sha256_of(result.local_path)
            done.add(key)
            stats["downloaded"] += 1
            stats["bytes"] += result.size_bytes
        else:
            stats["failed"] += 1
        append_manifest(record)
        mark = "OK " if result.status == "downloaded" else "FAIL"
        print(
            f"  {mark} {result.size_bytes:>9,}B  {link.filename[:70]}",
            flush=True,
        )


# --------------------------------------------------------------------------- KICPA


def kicpa_list_page(session, page: int) -> tuple[list[dict], int | None]:
    response = session.get(
        KICPA_LIST, params={"boardId": KICPA_BOARD, "page": page}, timeout=30
    )
    response.raise_for_status()
    soup = BeautifulSoup(response.content, "html.parser")
    items: list[dict] = []
    for anchor in soup.find_all("a"):
        match = KICPA_BLTN_RE.search(str(anchor.get("onclick", "")))
        if match is None:
            continue
        items.append(
            {"bltn_no": match.group(2), "title": clean_text(anchor.get_text(" ", strip=True))}
        )
    total_match = re.search(r"총\s*([\d,]+)\s*건", soup.get_text(" ", strip=True))
    total = int(total_match.group(1).replace(",", "")) if total_match else None
    return items, total


def collect_kicpa_item(
    session,
    item: dict,
    done: set[str],
    stats: dict,
    limit: int | None,
) -> None:
    url = f"{KICPA_READ}?boardId={KICPA_BOARD}&bltnNo={item['bltn_no']}"
    try:
        response = session.get(url, timeout=30)
    except Exception as exc:  # noqa: BLE001
        print(f"  [ERR] {url} :: {exc}", flush=True)
        stats["page_errors"] += 1
        return
    if response.status_code != 200:
        print(f"  [{response.status_code}] {url}", flush=True)
        stats["page_errors"] += 1
        return

    soup = BeautifulSoup(response.content, "html.parser")
    stats["pages"] += 1
    text = soup.get_text(" ", strip=True)
    issued = re.search(r"발행년월\s*([\d.]+)", text)
    kifrs = re.search(r"K-IFRS 번호\s*([\d,\s]+?)\s*(?:첨부파일|목록)", text)

    out_dir = OUT_ROOT / "kicpa"
    for anchor in soup.find_all("a", href=True):
        href = str(anchor.get("href", ""))
        if "fileMngr" not in href:
            continue
        seq = parse_qs(urlparse(href).query).get("fileSeq", ["1"])[0]
        key = f"kicpa:{item['bltn_no']}:{seq}"
        if key in done:
            stats["skipped"] += 1
            continue
        if limit is not None and stats["downloaded"] >= limit:
            return
        link = AttachmentLink(
            url=urljoin(url, href),
            filename=clean_text(anchor.get_text(" ", strip=True)) or f"{item['bltn_no']}.bin",
            method="GET",
        )
        result = download_attachment(session, link, out_dir, f"{item['bltn_no']}_{seq}", DELAY)
        record = {
            "key": key,
            "source": "kicpa",
            "section": "IFRS 실무사례",
            "title": item["title"],
            "page_url": url,
            "bltn_no": item["bltn_no"],
            "file_seq": seq,
            "issued": issued.group(1) if issued else None,
            "kifrs_no": clean_text(kifrs.group(1)) if kifrs else None,
            "filename": link.filename,
            "status": result.status,
            "content_type": result.content_type,
            "size_bytes": result.size_bytes,
        }
        if result.local_path is not None:
            record["path"] = str(result.local_path.relative_to(OUT_ROOT))
            record["sha256"] = sha256_of(result.local_path)
            done.add(key)
            stats["downloaded"] += 1
            stats["bytes"] += result.size_bytes
        else:
            stats["failed"] += 1
        append_manifest(record)
        mark = "OK " if result.status == "downloaded" else "FAIL"
        print(f"  {mark} {result.size_bytes:>9,}B  {link.filename[:70]}", flush=True)


# --------------------------------------------------------------------------- main


def main() -> int:
    parser = argparse.ArgumentParser(prog="collect_acct_data")
    _ = parser.add_argument("--phase", choices=["A", "B", "C"], action="append", default=None)
    _ = parser.add_argument("--limit", type=int, default=None, help="다운로드 파일 수 상한(스모크)")
    args = parser.parse_args()
    phases = set(args.phase or ["A", "B", "C"])

    OUT_ROOT.mkdir(parents=True, exist_ok=True)
    _, done = load_manifest()
    stats = {
        "pages": 0,
        "downloaded": 0,
        "skipped": 0,
        "failed": 0,
        "page_errors": 0,
        "bytes": 0,
    }
    started = time.time()
    print(f"[start] 기존 완료 {len(done)}건 / 출력 {OUT_ROOT}", flush=True)

    details, lists, kicpa_url = read_xlsx_links()
    session = new_session()

    if "A" in phases:
        print(f"\n=== Phase A: xlsx 상세 {len(details)}건 ===", flush=True)
        for index, entry in enumerate(details, 1):
            if args.limit is not None and stats["downloaded"] >= args.limit:
                break
            print(f"[A {index}/{len(details)}] {entry['title'][:60]}", flush=True)
            collect_kasb_page(session, entry, done, stats, args.limit)
            time.sleep(DELAY)

    if "B" in phases:
        print(f"\n=== Phase B: 목록 페이지 {len(lists)}건(A 미포함분) ===", flush=True)
        for index, entry in enumerate(lists, 1):
            if args.limit is not None and stats["downloaded"] >= args.limit:
                break
            print(f"[B {index}/{len(lists)}] {entry['section']} {entry['url']}", flush=True)
            collect_kasb_page(session, entry, done, stats, args.limit)
            time.sleep(DELAY)

    if "C" in phases and kicpa_url:
        print("\n=== Phase C: KICPA IFRS 실무사례 ===", flush=True)
        _ = session.get(kicpa_url, timeout=30)  # 세션(jsessionid) 확보
        page = 1
        total: int | None = None
        seen_bltn: set[str] = set()
        while True:
            if args.limit is not None and stats["downloaded"] >= args.limit:
                break
            items, page_total = kicpa_list_page(session, page)
            total = total or page_total
            fresh = [i for i in items if i["bltn_no"] not in seen_bltn]
            if not fresh:
                break
            for item in fresh:
                seen_bltn.add(item["bltn_no"])
                if args.limit is not None and stats["downloaded"] >= args.limit:
                    break
                print(f"[C p{page} {len(seen_bltn)}/{total or '?'}] {item['title'][:60]}", flush=True)
                collect_kicpa_item(session, item, done, stats, args.limit)
                time.sleep(DELAY)
            page += 1
            if total is not None and len(seen_bltn) >= total:
                break

    elapsed = time.time() - started
    print(
        f"\n[done] {elapsed / 60:.1f}분 | 페이지 {stats['pages']} "
        f"(오류 {stats['page_errors']}) | 다운로드 {stats['downloaded']} "
        f"| 스킵 {stats['skipped']} | 실패 {stats['failed']} "
        f"| {stats['bytes'] / 1_048_576:.1f}MB",
        flush=True,
    )
    return 0 if stats["failed"] == 0 and stats["page_errors"] == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
