# -*- coding: utf-8 -*-
"""Build a single CSV (and XLSX) for client delivery.

Reads `data/audit/coverage_report.csv` + `data/audit/scroll_index.json`
and produces:
- data/audit/client_report.csv       (UTF-8 with BOM for Excel)
- data/audit/client_report.xlsx      (auto-filter, frozen header)

Each of the 1,994 entries gets clear Korean-labelled columns describing
collection feasibility, blocking reason, recovery path, and recommended
action — so a non-developer can filter/sort directly in Excel.
"""
from __future__ import annotations

import csv
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
COVERAGE = ROOT / "data" / "audit" / "coverage_report.csv"
SCROLL = ROOT / "data" / "audit" / "scroll_index.json"
OUT_CSV = ROOT / "data" / "audit" / "client_report.csv"
OUT_XLSX = ROOT / "data" / "audit" / "client_report.xlsx"


# ---------------------------------------------------------------------------
# Classification logic
# ---------------------------------------------------------------------------

def classify(row: dict) -> dict:
    """Return Korean-labelled status fields for one coverage row."""
    rc = row.get("render_class", "")
    rs = row.get("recovery_status", "")
    br = row.get("block_reason", "")
    err = row.get("error", "") or ""
    a_status = row.get("analyzer_status", "")

    # Default
    tier = ""
    can_collect = ""
    reason = ""
    recovery = ""
    action = ""

    if rc == "robots_blocked":
        tier = "A. 영구불가"
        can_collect = "불가능"
        reason = "robots.txt 자동수집 거부"
        recovery = "사이트 운영자 협조시에만"
        action = "사이트 운영자에게 데이터 협조 요청 또는 제외"
    elif br == "cloudflare_403":
        tier = "A. 영구불가"
        can_collect = "불가능"
        reason = "Cloudflare IP 정책 차단 (403)"
        recovery = "Residential Proxy 도입 시 가능성"
        action = "Proxy 도입 검토 또는 제외"
    elif br in ("still_403",):
        tier = "A. 영구불가"
        can_collect = "불가능"
        reason = "서버 직접 403 거부 (Apache/Pantheon 정책)"
        recovery = "사이트 정책 변경시에만"
        action = "수동 등록 검토 또는 제외"
    elif br == "still_401":
        tier = "A. 영구불가"
        can_collect = "불가능"
        reason = "서버 401 인증 요구"
        recovery = "사이트 인증 협조시에만"
        action = "사이트 운영자에게 자격 요청"
    elif br == "other_404":
        tier = "A. 영구불가"
        can_collect = "불가능"
        reason = "URL 사라짐 (404)"
        recovery = "새 URL 제공시 가능"
        action = "클라이언트가 새 URL 제공 요청"
    elif rc == "dead":
        tier = "B. 응답불가"
        can_collect = "불가능"
        if "timeout" in err:
            reason = "사이트 응답시간 초과 (timeout)"
            recovery = "사이트 부하 해소시 일부 회복"
            action = "추후 재시도"
        elif "dns" in err:
            reason = "DNS 조회 실패 — 도메인 변경/소멸"
            recovery = "도메인 복구시에만"
            action = "URL 검증 / 새 URL 요청"
        elif "ssl" in err:
            reason = "TLS/SSL 인증서 영구 오류"
            recovery = "사이트 인증서 수정시"
            action = "사이트 운영자에게 통보"
        elif "connect_error" in err:
            reason = "호스트 연결 거부 또는 다운"
            recovery = "사이트 복구시"
            action = "추후 재시도"
        else:
            reason = f"HTTP 오류 또는 기타 (status={row.get('http_status','')})"
            recovery = "사이트 복구시"
            action = "추후 재시도"
    elif rc == "unknown":
        tier = "C. JS-only"
        can_collect = "불가능"
        reason = "사이트가 정적 HTML 미반환 (JS redirect 만 응답)"
        recovery = "Playwright 헤드리스 브라우저 도입 시 일부 가능"
        action = "Playwright 도입 검토"
    elif br == "cloudflare_challenge":
        tier = "D. Cloudflare 챌린지"
        can_collect = "현재불가 (보강시 가능)"
        reason = "Cloudflare JS 챌린지 — 일반 HTTP 클라이언트로 통과 불가"
        recovery = "Playwright 헤드리스 브라우저로 ~80% 회복 가능"
        action = "Playwright 분기 추가 권장"
    elif rc == "auth_blocked":
        # Anything still labelled auth_blocked but not classified above
        tier = "A. 영구불가"
        can_collect = "불가능"
        reason = "기타 인증 차단"
        recovery = "사이트 정책 의존"
        action = "수동 검토"
    else:
        # Collectable candidates
        if a_status == "ok":
            tier = "수집가능 (analyzer ok)"
            can_collect = "수집가능"
            reason = "정상 응답 + LLM 셀렉터 추론 성공"
            recovery = "—"
            action = "정상 배치 진입"
        elif a_status == "partial":
            tier = "수집가능 (partial)"
            can_collect = "수집가능"
            reason = "정상 응답 + 일부 셀렉터 추론"
            recovery = "—"
            action = "정상 배치 진입 (보조 보강 가능)"
        elif a_status == "fail":
            tier = "보강필요 (analyzer fail)"
            can_collect = "현재불가"
            reason = "LLM 셀렉터 추론 실패 (사이트 구조 특이 또는 다국어)"
            recovery = "다국어 프롬프트 / gpt-5.5 재시도 / 수동 셀렉터"
            action = "다국어 보강 또는 수동 검토"
        else:
            tier = "기타"
            can_collect = "확인필요"
            reason = "분류 미결"
            recovery = "—"
            action = "수동 확인"

    return {
        "수집분류": tier,
        "수집가능여부": can_collect,
        "사유": reason,
        "회복가능성": recovery,
        "권장조치": action,
    }


# ---------------------------------------------------------------------------
# Build CSV
# ---------------------------------------------------------------------------

def _load_tier1_results() -> dict:
    """Map entry_id → tier1_status from data/audit/tier1_runs/*.json."""
    result = {}
    tier1_dir = ROOT / "data" / "audit" / "tier1_runs"
    if not tier1_dir.exists():
        return result
    for f in tier1_dir.glob("*.json"):
        try:
            with f.open(encoding="utf-8") as fp:
                d = json.load(fp)
            result[d.get("entry_id", "")] = d.get("tier1_status", "")
        except Exception:
            continue
    return result


def _load_libertree_sites() -> set:
    """Set of host strings that have at least 1 document in libertree.db."""
    import sqlite3
    db = ROOT / "data" / "libertree.db"
    if not db.exists():
        return set()
    try:
        con = sqlite3.connect(f"file:{db}?mode=ro", uri=True)
        rows = con.execute(
            "SELECT DISTINCT s.site_url FROM sites s "
            "JOIN documents d ON d.site_id = s.site_id"
        ).fetchall()
        con.close()
        # Extract host from each site_url
        from urllib.parse import urlparse
        return {urlparse(r[0]).netloc.lower() for r in rows if r[0]}
    except Exception:
        return set()


def main() -> None:
    # Load coverage rows
    with open(COVERAGE, encoding="utf-8") as f:
        coverage_rows = list(csv.DictReader(f))

    # Load scroll_index for 기관명 + 일련번호 + 구분 + 세부구분
    with open(SCROLL, encoding="utf-8") as f:
        scroll_data = json.load(f)
    by_id = {e["entry_id"]: e for e in scroll_data["entries"]}

    # Tier 1 batch results + libertree.db actual collection
    tier1 = _load_tier1_results()
    libertree_hosts = _load_libertree_sites()

    # Korean column order
    cols = [
        "시트명",
        "일련번호",
        "기관명",
        "URL",
        "도메인",
        "콘텐츠유형",
        "세부구분",
        "수집분류",
        "수집가능여부",
        "사유",
        "회복가능성",
        "권장조치",
        "HTTP_상태",
        "robots_허용",
        "render_class",
        "block_reason",
        "analyzer_상태",
        "Tier1_결과",
        "실제수집됨",
        "비고_entry_id",
    ]

    out_rows = []
    for r in coverage_rows:
        eid = r["entry_id"]
        meta = by_id.get(eid, {})
        cls = classify(r)
        out_rows.append({
            "시트명": meta.get("sheet", ""),
            "일련번호": meta.get("일련번호", "") or "",
            "기관명": meta.get("name", "") or "",
            "URL": r.get("url", ""),
            "도메인": r.get("host", ""),
            "콘텐츠유형": meta.get("gubun", "") or "",
            "세부구분": meta.get("detail", "") or "",
            "수집분류": cls["수집분류"],
            "수집가능여부": cls["수집가능여부"],
            "사유": cls["사유"],
            "회복가능성": cls["회복가능성"],
            "권장조치": cls["권장조치"],
            "HTTP_상태": r.get("http_status", ""),
            "robots_허용": "예" if r.get("robots_ok") == "true" else ("아니오" if r.get("robots_ok") == "false" else ""),
            "render_class": r.get("render_class", ""),
            "block_reason": r.get("block_reason", ""),
            "analyzer_상태": r.get("analyzer_status", ""),
            "Tier1_결과": tier1.get(eid, ""),
            "실제수집됨": "예" if (r.get("host", "") or "").lower() in libertree_hosts else "",
            "비고_entry_id": eid,
        })

    # Sort: 불가/보강필요 먼저, 가능 나중
    sort_priority = {
        "A. 영구불가": 0,
        "B. 응답불가": 1,
        "C. JS-only": 2,
        "D. Cloudflare 챌린지": 3,
        "보강필요 (analyzer fail)": 4,
        "수집가능 (partial)": 5,
        "수집가능 (analyzer ok)": 6,
        "기타": 7,
    }
    out_rows.sort(key=lambda r: (sort_priority.get(r["수집분류"], 99), r["시트명"], r["도메인"]))

    # Write CSV with BOM (Excel friendly for Korean)
    with open(OUT_CSV, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(out_rows)

    print(f"[client-report] wrote {OUT_CSV} ({len(out_rows)} rows)")

    # Build XLSX with auto-filter and frozen header
    try:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font, Alignment
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "수집가능성 평가"

        # Header style
        header_fill = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.append(cols)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align

        # Conditional row styling per 수집분류
        tier_fills = {
            "A. 영구불가": PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"),
            "B. 응답불가": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
            "C. JS-only": PatternFill(start_color="DEEBF7", end_color="DEEBF7", fill_type="solid"),
            "D. Cloudflare 챌린지": PatternFill(start_color="DEEBF7", end_color="DEEBF7", fill_type="solid"),
            "보강필요 (analyzer fail)": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
            "수집가능 (partial)": PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
            "수집가능 (analyzer ok)": PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
        }

        for row in out_rows:
            ws.append([row[c] for c in cols])
            r_idx = ws.max_row
            fill = tier_fills.get(row["수집분류"])
            if fill:
                for cell in ws[r_idx]:
                    cell.fill = fill

        # Column widths (rough)
        widths = {
            "시트명": 18,
            "일련번호": 8,
            "기관명": 35,
            "URL": 60,
            "도메인": 28,
            "콘텐츠유형": 18,
            "세부구분": 18,
            "수집분류": 22,
            "수집가능여부": 18,
            "사유": 40,
            "회복가능성": 35,
            "권장조치": 35,
            "HTTP_상태": 10,
            "robots_허용": 10,
            "render_class": 15,
            "block_reason": 22,
            "analyzer_상태": 13,
            "Tier1_결과": 14,
            "실제수집됨": 12,
            "비고_entry_id": 13,
        }
        for i, c in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 14)

        # Freeze header + auto filter
        ws.freeze_panes = "A2"
        last_col = get_column_letter(len(cols))
        ws.auto_filter.ref = f"A1:{last_col}{ws.max_row}"

        # Summary sheet
        ws2 = wb.create_sheet("요약")
        ws2.append(["수집분류", "건수", "비율 (%)", "설명"])
        for cell in ws2[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align

        from collections import Counter
        cnt = Counter(r["수집분류"] for r in out_rows)
        total = len(out_rows)
        descs = {
            "A. 영구불가": "사이트가 정책·차단으로 자동 수집 거부 (robots.txt, IP 차단, 서버 403)",
            "B. 응답불가": "사이트 자체가 응답하지 않음 (timeout, DNS, SSL 등)",
            "C. JS-only": "정적 HTML 미반환 (JS redirect)",
            "D. Cloudflare 챌린지": "JS 챌린지 통과 필요 (Playwright 도입 시 회복 가능)",
            "보강필요 (analyzer fail)": "LLM 셀렉터 추론 실패 (모델 보강 시 회복 가능)",
            "수집가능 (analyzer ok)": "정상 응답 + 셀렉터 추론 성공 — 즉시 수집 가능",
            "수집가능 (partial)": "셀렉터 추론 일부 — 보조 보강시 가능",
            "기타": "분류 미결",
        }
        order = ["A. 영구불가", "B. 응답불가", "C. JS-only", "D. Cloudflare 챌린지",
                 "보강필요 (analyzer fail)", "수집가능 (partial)", "수집가능 (analyzer ok)", "기타"]
        for k in order:
            v = cnt.get(k, 0)
            if v == 0:
                continue
            pct = round(100 * v / total, 1)
            ws2.append([k, v, pct, descs.get(k, "")])
            r_idx = ws2.max_row
            fill = tier_fills.get(k)
            if fill:
                for cell in ws2[r_idx]:
                    cell.fill = fill

        # Total row
        ws2.append(["합계", total, 100.0, ""])
        for cell in ws2[ws2.max_row]:
            cell.font = Font(bold=True)

        ws2.column_dimensions["A"].width = 25
        ws2.column_dimensions["B"].width = 10
        ws2.column_dimensions["C"].width = 12
        ws2.column_dimensions["D"].width = 60
        ws2.freeze_panes = "A2"

        # Move 요약 to first
        wb.move_sheet(ws2, offset=-1)

        wb.save(OUT_XLSX)
        print(f"[client-report] wrote {OUT_XLSX}")
    except ImportError:
        print("[client-report] openpyxl not installed — skipped XLSX")


if __name__ == "__main__":
    main()
