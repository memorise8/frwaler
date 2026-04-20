#!/usr/bin/env python3
"""Fino Search - batch search input.xlsx entries against taxlaw.nts.go.kr.

Strategy:
  Pass 0: document-mode search via ASEISA001MR01 (searchType="document")
          — searches only doc-number fields (NTST_DCM_DSCM_CNTN, DOCU_NO_STR1/2/3)
          with automatic 0-prefix padding variants. Accurate, low-noise.
  Pass 1 (fallback): viewCount=50 keyword search (icldVcbCtl) with exact match filter.
  Pass 2 (fallback): viewCount=200 + pagination (5 pages).
  Pass 3 (fallback): viewCount=500 + pagination (10 pages).

Outputs:
  docs/fino-seach-results.md     -- matched entries table
  docs/fino-seach-unmatched.md   -- unmatched entries
  .cache/fino_search.json        -- resume checkpoint
"""

from __future__ import annotations

import argparse
import json
import os
import re
import subprocess
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
INPUT_XLSX = ROOT / "input.xlsx"
RESULTS_MD = ROOT / "docs" / "fino-seach-results.md"
UNMATCHED_MD = ROOT / "docs" / "fino-seach-unmatched.md"
CACHE_DIR = ROOT / ".cache"
CHECKPOINT = CACHE_DIR / "fino_search.json"

API_URL = "https://taxlaw.nts.go.kr/action.do"
REFERER = "https://taxlaw.nts.go.kr"

COLLECTIONS = [
    ("판례", "precedent,precedent_gr",
     ["001_05", "001_06", "001_07", "001_08", "001_09", "001_10"]),
    ("해석례", "question,question_gr",
     ["001_01", "001_02", "001_03", "001_04"]),
]

TAG_RE = re.compile(r"<!H[SE]>")


def strip_tags(s: str) -> str:
    return TAG_RE.sub("", s or "")


def normalize_input(raw: str) -> str:
    """Strip .md suffix only; preserve trailing hyphens (e.g. 헌법재판소 `2000-헌마-8-`)."""
    s = str(raw).strip()
    if s.lower().endswith(".md"):
        s = s[:-3]
    return s.strip()


def read_inputs(path: Path) -> list[str]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [r[0] for r in ws.iter_rows(values_only=True) if r[0]]
    seen: set[str] = set()
    uniq: list[str] = []
    for r in rows[1:]:  # skip header
        n = normalize_input(r)
        if n and n not in seen:
            seen.add(n)
            uniq.append(n)
    return uniq


def api_call(keyword: str, collection: str, codes: list[str],
             view_count: int = 50, start: int = 1,
             date_from: str | None = None,
             date_to: str | None = None,
             timeout: int = 35) -> list[dict]:
    params = {
        "collectionName": collection,
        "sortField": "DCM_RGT_DTM/DESC",
        "startCount": start,
        "viewCount": view_count,
        "dcmClCdCtl": codes,
        "qstnPrdcOrgnClCtl": [],
        "rltnStttCtl": [],
        "schDtBase": "DCM_RGT_DTM",
        "icldVcbCtl": [keyword],
    }
    if date_from:
        params["bltnStrtDt"] = date_from
    if date_to:
        params["bltnEndDt"] = date_to
    param_json = json.dumps(params, ensure_ascii=False)
    cmd = [
        "curl", "-skL", "--max-time", "30", "-X", "POST",
        "-H", "Content-Type: application/x-www-form-urlencoded",
        "-H", f"Referer: {REFERER}",
        "--data-urlencode", f"paramData={param_json}",
        "-d", "actionId=ASIPDI002PR01",
        API_URL,
    ]
    for attempt in range(3):
        try:
            r = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
            if not r.stdout.strip():
                time.sleep(1 + attempt)
                continue
            data = json.loads(r.stdout)
            return data.get("data", {}).get("ASIPDI002PR01", {}).get("body", []) or []
        except (subprocess.TimeoutExpired, json.JSONDecodeError):
            time.sleep(1 + attempt)
    return []


def find_exact_match(keyword: str, body: list[dict]) -> dict | None:
    for item in body:
        dcm = item.get("dcm", {}) or {}
        doc_no = strip_tags(dcm.get("NTST_DCM_DSCM_CNTN", ""))
        if doc_no == keyword:
            return dcm
    return None


def extract_fields(keyword: str, dcm: dict) -> dict:
    rgt = dcm.get("DCM_RGT_DTM", "") or ""
    if len(rgt) >= 8:
        pub = f"{rgt[0:4]}-{rgt[4:6]}-{rgt[6:8]}"
    else:
        pub = ""
    return {
        "keyword": keyword,
        "title": strip_tags(dcm.get("TTL", "")),
        "type": dcm.get("NTST_DCM_CL_NM", "") or "",
        "category": dcm.get("NTST_TLAW_CL_NM", "") or "",
        "published": pub,
        "doc_id": dcm.get("DOC_ID", ""),
    }


DOC_MODE_COLLECTIONS = [("판례", "precedent"), ("해석례", "question")]


def doc_mode_api_call(keyword: str, collection: str,
                      view_count: int = 10,
                      timeout: int = 35) -> list[dict]:
    """Document-mode search via ASEISA001MR01.

    Server searches only doc-number indexes (NTST_DCM_DSCM_CNTN, DOCU_NO_STR1/2/3)
    with automatic 0-prefix padding variants. Returns flat items (not nested
    under `dcm` key).
    """
    params = {
        "schVcb": keyword,
        "startCount": 1,
        "collection": collection,
        "wnKey": "",
        "searchType": "document",
        "sortField": "SCORE/DESC",
        "ntstTlawClCdList": [],
        "icldVcbCtl": [],
        "exclVcbCtl": [],
        "rltnStttCtl": [],
        "schDtBase": "DCM_RGT_DTM",
        "viewCount": str(view_count),
        "prtsSprcChiefJdgmYn": "",
        "prtsAttrYrCtl": [],
        "prtsPrgrStatCtl": [],
        "prtsLwsDfntYn": "",
        "infrOpClCtl": [],
        "mainIdCtl": [],
        "useSynonymYn": "N",
    }
    cmd = [
        "curl", "-skL", "--max-time", "30", "-X", "POST",
        "-H", "Content-Type: application/x-www-form-urlencoded; charset=UTF-8",
        "-H", "Origin: https://taxlaw.nts.go.kr",
        "-H", "Referer: https://taxlaw.nts.go.kr/is/USEISA001M.do",
        "-H", "X-Requested-With: XMLHttpRequest",
        "--data-urlencode", f"paramData={json.dumps(params, ensure_ascii=False)}",
        "-d", "actionId=ASEISA001MR01",
        API_URL,
    ]
    for attempt in range(3):
        try:
            r = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
            if not r.stdout.strip():
                time.sleep(1 + attempt)
                continue
            data = json.loads(r.stdout)
            items: list[dict] = []
            coll_list = (data.get("data", {}).get("ASEISA001MR01", {})
                         .get("searchResultVO", {}).get("collectionList", []) or [])
            for c in coll_list:
                if c and c.get("resultList"):
                    items.extend(c["resultList"])
            return items
        except (subprocess.TimeoutExpired, json.JSONDecodeError):
            time.sleep(1 + attempt)
    return []


def find_exact_match_flat(keyword: str, items: list[dict]) -> dict | None:
    """Exact match filter for flat item structure (document mode response)."""
    for item in items:
        doc_no = strip_tags(item.get("NTST_DCM_DSCM_CNTN", ""))
        if doc_no == keyword:
            return item
    return None


def search_pass0(keyword: str) -> dict | None:
    """Document-mode search — primary strategy (accurate, low-noise)."""
    for label, coll in DOC_MODE_COLLECTIONS:
        items = doc_mode_api_call(keyword, coll, view_count=10)
        dcm = find_exact_match_flat(keyword, items)
        if dcm:
            r = extract_fields(keyword, dcm)
            r["source"] = label
            return r
    return None


def search_pass1(keyword: str) -> dict | None:
    """Single viewCount=50 search across both collections."""
    for label, coll, codes in COLLECTIONS:
        body = api_call(keyword, coll, codes, view_count=50)
        dcm = find_exact_match(keyword, body)
        if dcm:
            r = extract_fields(keyword, dcm)
            r["source"] = label
            return r
    return None


def search_pass2(keyword: str, max_pages: int = 5,
                 view_count: int = 200) -> dict | None:
    """Pagination-based retry for unmatched entries."""
    for label, coll, codes in COLLECTIONS:
        for page in range(max_pages):
            start = page * view_count + 1
            body = api_call(keyword, coll, codes,
                            view_count=view_count, start=start)
            if not body:
                break
            dcm = find_exact_match(keyword, body)
            if dcm:
                r = extract_fields(keyword, dcm)
                r["source"] = label
                return r
            if len(body) < view_count:
                break
            time.sleep(0.2)
    return None


def search_pass3(keyword: str, max_pages: int = 10,
                 view_count: int = 500) -> dict | None:
    """Exhaustive retry with larger page size."""
    return search_pass2(keyword, max_pages=max_pages, view_count=view_count)


def load_checkpoint() -> dict:
    if CHECKPOINT.exists():
        try:
            return json.loads(CHECKPOINT.read_text(encoding="utf-8"))
        except Exception:
            pass
    return {"matched": {}, "unmatched": [], "pass1_done": False}


def save_checkpoint(state: dict) -> None:
    CACHE_DIR.mkdir(exist_ok=True)
    tmp = CHECKPOINT.with_suffix(".tmp")
    tmp.write_text(json.dumps(state, ensure_ascii=False, indent=1),
                   encoding="utf-8")
    tmp.replace(CHECKPOINT)


def md_escape(s: str) -> str:
    return (s or "").replace("|", "\\|").replace("\n", " ").strip()


def write_outputs(state: dict) -> None:
    matched = state["matched"]
    unmatched = state["unmatched"]

    lines = ["# Fino Search 결과", "",
             f"`input.xlsx` 문서번호를 `taxlaw.nts.go.kr`에서 검색한 결과.",
             "",
             f"- 매치: **{len(matched)}건**",
             f"- 미매치: **{len(unmatched)}건**",
             "",
             "| 검색어 | 제목 | 유형 | 분류 | 생산일자 |",
             "|--------|------|------|------|----------|"]
    for kw in sorted(matched.keys()):
        r = matched[kw]
        lines.append(f"| {md_escape(r['keyword'])} | {md_escape(r['title'])} | "
                     f"{md_escape(r['type'])} | {md_escape(r['category'])} | "
                     f"{md_escape(r['published'])} |")
    RESULTS_MD.write_text("\n".join(lines) + "\n", encoding="utf-8")

    u_lines = ["# Fino Search 미매치 목록", "",
               f"정확 매치되지 않은 {len(unmatched)}건.",
               "",
               "| # | 검색어 |",
               "|---|--------|"]
    for i, kw in enumerate(unmatched, 1):
        u_lines.append(f"| {i} | {md_escape(kw)} |")
    UNMATCHED_MD.write_text("\n".join(u_lines) + "\n", encoding="utf-8")


def run_pass(keywords: list[str], state: dict, search_fn,
             workers: int = 3, delay: float = 0.0,
             save_every: int = 100, pass_name: str = "pass1",
             batch_size: int = 200) -> None:
    matched = state["matched"]
    remaining = [k for k in keywords if k not in matched]
    total = len(remaining)
    print(f"[{pass_name}] 대상 {total}건 (이미 매치 {len(matched)}건 스킵)",
          flush=True)
    if total == 0:
        return

    def worker(kw: str):
        if delay > 0:
            time.sleep(delay)
        return kw, search_fn(kw)

    done = 0
    last_save = time.time()
    start_t = time.time()
    with ThreadPoolExecutor(max_workers=workers) as ex:
        # submit in rolling batches to avoid creating 17k futures at once
        idx = 0
        pending = set()
        while idx < len(remaining) or pending:
            while len(pending) < batch_size and idx < len(remaining):
                pending.add(ex.submit(worker, remaining[idx]))
                idx += 1
            if not pending:
                break
            # wait for any completion
            done_set = {f for f in pending if f.done()}
            if not done_set:
                # block on one
                for f in as_completed(pending, timeout=None):
                    done_set = {f}
                    break
            for f in done_set:
                pending.discard(f)
                try:
                    kw, res = f.result()
                except Exception as e:
                    print(f"[{pass_name}] ERR: {e}", flush=True)
                    continue
                done += 1
                if res:
                    matched[kw] = res
                if done % 50 == 0 or done == total:
                    rate = done / max(1e-9, time.time() - start_t)
                    eta = (total - done) / max(1e-9, rate)
                    print(f"[{pass_name}] {done}/{total} "
                          f"매치 {len(matched)} "
                          f"({rate:.1f}/s ETA {eta/60:.1f}min)",
                          flush=True)
                if time.time() - last_save > 30 or done % save_every == 0:
                    save_checkpoint(state)
                    write_outputs(state)
                    last_save = time.time()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--limit", type=int, default=None,
                    help="max keywords to process (for testing)")
    ap.add_argument("--workers", type=int, default=3)
    ap.add_argument("--delay", type=float, default=0.3,
                    help="inter-request submission delay (per worker)")
    ap.add_argument("--skip-pass2", action="store_true")
    ap.add_argument("--only-pass2", action="store_true")
    ap.add_argument("--skip-pass0", action="store_true",
                    help="skip document-mode search (use legacy keyword passes only)")
    ap.add_argument("--skip-pass3", action="store_true")
    ap.add_argument("--reset", action="store_true")
    args = ap.parse_args()

    CACHE_DIR.mkdir(exist_ok=True)
    if args.reset and CHECKPOINT.exists():
        CHECKPOINT.unlink()

    keywords = read_inputs(INPUT_XLSX)
    if args.limit:
        keywords = keywords[:args.limit]
    print(f"총 고유 검색어: {len(keywords)}", flush=True)

    state = load_checkpoint()

    # Pass 0: document-mode (primary, accurate)
    if not args.skip_pass0 and not args.only_pass2:
        run_pass(keywords, state, search_pass0,
                 workers=args.workers, delay=args.delay,
                 pass_name="pass0-docmode")
        unmatched = [k for k in keywords if k not in state["matched"]]
        state["unmatched"] = unmatched
        save_checkpoint(state)
        write_outputs(state)
        print(f"\npass0 완료: 매치 {len(state['matched'])}, "
              f"미매치 {len(unmatched)}", flush=True)

    # Pass 1: keyword search (fallback)
    if not args.only_pass2:
        run_pass(keywords, state, search_pass1,
                 workers=args.workers, delay=args.delay,
                 pass_name="pass1")
        state["pass1_done"] = True
        unmatched = [k for k in keywords if k not in state["matched"]]
        state["unmatched"] = unmatched
        save_checkpoint(state)
        write_outputs(state)
        print(f"\npass1 완료: 매치 {len(state['matched'])}, "
              f"미매치 {len(unmatched)}", flush=True)

    # Pass 2: paginated keyword search (fallback)
    if not args.skip_pass2:
        unmatched = [k for k in keywords if k not in state["matched"]]
        print(f"\npass2 시작: {len(unmatched)}건 재검색 "
              f"(viewCount=200, 최대 5페이지)", flush=True)
        run_pass(unmatched, state, search_pass2,
                 workers=max(1, args.workers - 1),
                 delay=args.delay * 2,
                 pass_name="pass2")
        state["unmatched"] = [k for k in keywords if k not in state["matched"]]
        save_checkpoint(state)
        write_outputs(state)

    # Pass 3: deep paginated (fallback)
    if not args.skip_pass3:
        unmatched = [k for k in keywords if k not in state["matched"]]
        if unmatched:
            print(f"\npass3 시작: {len(unmatched)}건 "
                  f"(viewCount=500, 최대 10페이지)", flush=True)
            run_pass(unmatched, state, search_pass3,
                     workers=max(1, args.workers - 1),
                     delay=args.delay * 2,
                     pass_name="pass3")
            state["unmatched"] = [k for k in keywords if k not in state["matched"]]
            save_checkpoint(state)
            write_outputs(state)

    print(f"\n=== 최종 결과 ===", flush=True)
    print(f"매치: {len(state['matched'])}", flush=True)
    print(f"미매치: {len(state['unmatched'])}", flush=True)
    print(f"결과: {RESULTS_MD}", flush=True)
    print(f"미매치: {UNMATCHED_MD}", flush=True)


if __name__ == "__main__":
    main()
